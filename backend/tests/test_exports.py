"""Выгрузка таблиц в Excel (.xlsx): доступ super_admin, валидный workbook, нужные листы."""
import io

from openpyxl import load_workbook

from .conftest import make_order

# путь -> ожидаемые листы (в порядке создания)
EXPORTS = {
    "/api/admin/exports/orders.xlsx": ["Orders", "Items"],
    "/api/admin/exports/customers.xlsx": ["Customers"],
    "/api/admin/exports/payments.xlsx": ["Payments"],
    "/api/admin/exports/audience.xlsx": ["Overview", "Customers", "Segments", "Personas", "RFM grid"],
    "/api/admin/exports/dashboard.xlsx": ["Summary", "Daily", "Top products", "Top customers",
                                          "Size mix", "Top add-ons", "Add-on affinity", "By hour"],
}


def test_exports_return_valid_xlsx_with_expected_sheets(client, admin, customer):
    make_order(client, customer)  # гарантируем наличие данных
    for path, sheets in EXPORTS.items():
        r = client.get(path, headers=admin["headers"])
        assert r.status_code == 200, (path, r.text)
        assert "spreadsheetml.sheet" in r.headers["content-type"]
        assert ".xlsx" in r.headers["content-disposition"]
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == sheets, (path, wb.sheetnames)


def test_exports_require_super_admin(client, manager, customer):
    for path in EXPORTS:
        assert client.get(path).status_code == 401                         # без авторизации
        assert client.get(path, headers=manager["headers"]).status_code == 403   # manager не владелец
        assert client.get(path, headers=customer["headers"]).status_code == 403  # клиент — не staff


def test_orders_export_has_header_and_rows(client, admin, customer):
    make_order(client, customer)
    r = client.get("/api/admin/exports/orders.xlsx", headers=admin["headers"])
    ws = load_workbook(io.BytesIO(r.content))["Orders"]
    assert ws.cell(1, 1).value == "ID" and ws.cell(1, 16).value == "Total"
    assert ws.max_row >= 2  # шапка + хотя бы один заказ


def test_dashboard_daily_breakdown_zero_filled_over_period(client, admin):
    # фиксированный месяц -> 30 строк по дням (заполнение нулями), детерминированно
    r = client.get("/api/admin/exports/dashboard.xlsx?from=2026-06-01T00:00:00&to=2026-06-30T23:59:59",
                   headers=admin["headers"])
    assert r.status_code == 200
    daily = load_workbook(io.BytesIO(r.content))["Daily"]
    assert daily.cell(1, 1).value == "Date"
    assert daily.max_row == 1 + 30  # шапка + каждый день июня


def test_audience_export_customer_sheet_has_rfm_columns(client, admin, customer):
    make_order(client, customer)
    r = client.get("/api/admin/exports/audience.xlsx", headers=admin["headers"])
    ws = load_workbook(io.BytesIO(r.content))["Customers"]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    for col in ("Segment", "R", "F", "M", "Predicted CLV", "Personas"):
        assert col in headers, (col, headers)
