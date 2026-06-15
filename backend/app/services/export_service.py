"""Выгрузка таблиц в Excel (.xlsx) — заказы, клиенты, платежи, аудитория, дашборд.

Чистый слой: на вход подаются уже сериализованные данные (dict/list из роутеров и crm),
на выход — bytes готового .xlsx. Никаких обращений к БД и FastAPI здесь нет.

Каждая «книга» (Workbook) может содержать несколько листов; шапка выделена, верхняя строка
заморожена, ширина колонок автоподбирается, ISO-даты превращаются в настоящие даты Excel.
"""
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADER_FILL = PatternFill("solid", fgColor="4A56E2")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_DATE_FMT = "yyyy-mm-dd hh:mm"


def _coerce(v):
    """ISO-8601 строку превращаем в datetime (станет настоящей датой Excel); прочее — как есть."""
    if isinstance(v, str) and len(v) >= 10 and v[4:5] == "-" and v[7:8] == "-":
        try:
            return datetime.fromisoformat(v.replace("Z", "+00:00")).replace(tzinfo=None)
        except ValueError:
            return v
    return v


def _add_sheet(wb: Workbook, title: str, headers: list[str], rows: list[list]):
    """Лист с жирной шапкой, заморозкой строки и автошириной колонок."""
    ws = wb.create_sheet(title[:31])  # имя листа в Excel ≤ 31 символа
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(1, c)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    widths = [len(str(h)) for h in headers]
    for row in rows:
        values = [_coerce(v) for v in row]
        ws.append(values)
        r = ws.max_row
        for c, v in enumerate(values, start=1):
            if isinstance(v, datetime):
                ws.cell(r, c).number_format = _DATE_FMT
            shown = 16 if isinstance(v, datetime) else len(str("" if v is None else v))
            if c - 1 < len(widths):
                widths[c - 1] = max(widths[c - 1], shown)
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = min(max(w + 2, 10), 52)
    ws.freeze_panes = "A2"
    return ws


def _kv_sheet(wb: Workbook, title: str, pairs: list[tuple[str, object]]):
    """Лист «Метрика | Значение» для сводок/KPI."""
    return _add_sheet(wb, title, ["Metric", "Value"], [[k, v] for k, v in pairs])


def _new_wb() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)  # убираем дефолтный пустой лист — всегда создаём именованные
    return wb


def _save(wb: Workbook) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------- 1. Заказы ----------------

def build_orders_wb(rows: list[dict]) -> bytes:
    """Лист Orders (одна строка на заказ) + Items (одна строка на позицию с добавками)."""
    wb = _new_wb()
    order_headers = ["ID", "Number", "Status", "Payment", "Created", "Outlet", "Customer",
                     "Phone", "Car plate", "Emirate", "Manager ID", "Items", "Drinks",
                     "Subtotal", "Coupon discount", "Total", "Rating", "Arrived"]
    order_rows = []
    item_rows = []
    for o in rows:
        outlet = (o.get("outlet") or {}).get("name") if o.get("outlet") else None
        items = o.get("items") or []
        drinks = sum(i.get("quantity", 0) for i in items)
        order_rows.append([
            o["id"], o.get("number"), o.get("status"), o.get("paymentStatus"),
            o.get("createdAt"), outlet, o.get("customerName"), o.get("phone"),
            o.get("carPlate"), o.get("emirate"), o.get("managerId"), len(items), drinks,
            o.get("subtotal"), o.get("couponDiscount"), o.get("total"),
            o.get("rating"), "yes" if o.get("arrived") else "no",
        ])
        for i in items:
            addons = "; ".join(
                f"{a.get('nameEn') or a.get('name')} x{a.get('portions')}"
                + (f" ({a.get('amount')}{a.get('unit')})" if a.get("amount") else "")
                for a in (i.get("addons") or [])
            )
            item_rows.append([
                o.get("number"), o.get("createdAt"),
                i.get("drinkNameEn") or i.get("drinkName"), i.get("customName"),
                i.get("sizeLabel"), i.get("quantity"), i.get("unitPrice"),
                "yes" if i.get("paidByCoupon") else "", addons,
            ])
    _add_sheet(wb, "Orders", order_headers, order_rows)
    _add_sheet(wb, "Items", ["Order #", "Created", "Drink", "Custom name", "Size",
                             "Qty", "Unit price", "By coupon", "Add-ons"], item_rows)
    return _save(wb)


# ---------------- 2. Клиенты ----------------

def build_customers_wb(rows: list[dict]) -> bytes:
    headers = ["ID", "Name", "Phone", "Car plate", "Language", "Registered",
               "Paid orders", "Total spent", "Last order"]
    data = [[r["id"], r.get("name"), r.get("phone"), r.get("carPlate"), r.get("locale"),
             r.get("createdAt"), r.get("orders"), r.get("spent"), r.get("lastOrderAt")]
            for r in rows]
    wb = _new_wb()
    _add_sheet(wb, "Customers", headers, data)
    return _save(wb)


# ---------------- 3. Платежи ----------------

def build_payments_wb(rows: list[dict]) -> bytes:
    headers = ["ID", "Order #", "Created", "Customer", "Phone", "Status", "Method",
               "Card brand", "Last4", "Amount", "Fee", "Net", "Refunded", "Currency",
               "Provider", "Provider ID", "Risk", "Dispute", "Livemode"]
    data = [[
        p["id"], p.get("orderNumber"), p.get("createdAt"), p.get("customerName"),
        p.get("customerPhone"), p.get("status"), p.get("method"), p.get("cardBrand"),
        p.get("cardLast4"), p.get("amount"), p.get("fee"), p.get("net"), p.get("refunded"),
        p.get("currency"), p.get("provider"), p.get("providerId"), p.get("riskLevel"),
        p.get("disputeStatus"), "yes" if p.get("livemode") else "no",
    ] for p in rows]
    wb = _new_wb()
    _add_sheet(wb, "Payments", headers, data)
    return _save(wb)


# ---------------- 4. Аудитория (RFM / сегменты / персоны) ----------------

def build_audience_wb(audience: dict, customer_rows: list[dict]) -> bytes:
    """Многолистовая книга: Overview (KPI) · Customers (на клиента: сегмент/RFM/отток/CLV/
    персоны) · Segments · Personas · RFM grid."""
    wb = _new_wb()
    k = audience.get("kpis", {})
    _kv_sheet(wb, "Overview", [
        ("Customers total", audience.get("total")),
        ("Active (≤30d)", k.get("active")),
        ("At risk", k.get("atRisk")),
        ("Churned (>60d)", k.get("churned")),
        ("New this month", k.get("newThisMonth")),
        ("Avg predicted CLV", k.get("avgCLV")),
    ])

    # на каждого клиента — сегмент + RFM + отток + CLV + персоны (разделение по поведению)
    _add_sheet(wb, "Customers", [
        "ID", "Name", "Phone", "Segment", "R", "F", "M", "RFM", "Recency days",
        "Paid orders", "Total spent", "Avg order value", "Orders/month",
        "Churn risk", "Churn probability", "Predicted CLV", "Personas",
    ], [[
        c["id"], c.get("name"), c.get("phone"), c.get("segmentLabel"),
        c.get("r"), c.get("f"), c.get("m"), c.get("rfmScore"), c.get("recencyDays"),
        c.get("paidOrders"), c.get("totalSpent"), c.get("avgOrderValue"),
        c.get("ordersPerMonth"), c.get("churnRisk"), c.get("churnProbability"),
        c.get("clvPredicted"), ", ".join(c.get("personas") or []),
    ] for c in customer_rows])

    _add_sheet(wb, "Segments", [
        "Key", "Segment", "Description", "Customers", "Share", "Avg spent",
        "Avg recency", "Avg frequency", "Top personas",
    ], [[
        s.get("key"), s.get("label"), s.get("description"), s.get("count"), s.get("share"),
        s.get("avgSpent"), s.get("avgRecency"), s.get("avgFrequency"),
        ", ".join(f"{t['tag']} ({t['count']})" for t in (s.get("personaTags") or [])),
    ] for s in audience.get("segments", [])])

    _add_sheet(wb, "Personas", ["Persona", "Customers", "Share"],
               [[p.get("tag"), p.get("count"), p.get("share")]
                for p in audience.get("personas", [])])

    # RFM-сетка 5×5: строки Frequency 5→1, столбцы Recency 1→5 (как на дашборде)
    grid = audience.get("rfmGrid") or [[0] * 5 for _ in range(5)]
    rfm_rows = []
    for f in range(5, 0, -1):
        rfm_rows.append([f"F{f}"] + [grid[f - 1][r - 1] for r in range(1, 6)])
    _add_sheet(wb, "RFM grid", ["Frequency \\ Recency", "R1", "R2", "R3", "R4", "R5"], rfm_rows)
    return _save(wb)


# ---------------- 5. Дашборд (сводка + по дням + разбивки) ----------------

def build_dashboard_wb(summary: dict, daily: list[dict], period: dict) -> bytes:
    """Многолистовая книга: Summary (KPI периода + дельты + время сервиса) · Daily (метрики по
    дням) · Top products · Top customers · Size mix · Top add-ons · Add-on affinity · By hour."""
    wb = _new_wb()
    deltas = summary.get("deltas") or {}
    st = summary.get("serviceTime") or {}

    def _d(key):
        v = deltas.get(key)
        return f"{v:+.1f}%" if v is not None else "—"

    _kv_sheet(wb, "Summary", [
        ("Period from", period.get("from") or "all time"),
        ("Period to", period.get("to") or "now"),
        ("Outlet", period.get("outlet") or "all outlets"),
        ("Revenue (AED)", summary.get("revenue")),
        ("Revenue Δ vs prev", _d("revenue")),
        ("Orders", summary.get("ordersCount")),
        ("Orders Δ vs prev", _d("ordersCount")),
        ("Drinks sold", summary.get("drinksSold")),
        ("Drinks Δ vs prev", _d("drinksSold")),
        ("Avg order value", summary.get("avgOrderValue")),
        ("AOV Δ vs prev", _d("avgOrderValue")),
        ("Avg drinks/order", summary.get("avgDrinksPerOrder")),
        ("Avg add-ons/drink", summary.get("avgAddons")),
        ("Prep time (min)", st.get("prepMin")),
        ("Pickup time (min)", st.get("pickupMin")),
        ("Total service (min)", st.get("totalMin")),
        ("Service samples", st.get("samples")),
    ])

    _add_sheet(wb, "Daily", ["Date", "Orders", "Revenue", "Drinks", "Avg order value",
                             "New customers"],
               [[d["date"], d["orders"], d["revenue"], d["drinks"],
                 d["avgOrderValue"], d["newCustomers"]] for d in daily])

    _add_sheet(wb, "Top products", ["Product", "Slug", "Revenue", "Qty"],
               [[p.get("name"), p.get("slug"), p.get("revenue"), p.get("qty")]
                for p in summary.get("topProducts", [])])

    _add_sheet(wb, "Top customers", ["User ID", "Name", "Phone", "Orders", "Spent", "Last order"],
               [[c.get("userId"), c.get("name"), c.get("phone"), c.get("orders"),
                 c.get("spent"), c.get("lastOrderAt")]
                for c in summary.get("topCustomers", [])])

    _add_sheet(wb, "Size mix", ["Size", "Qty", "Share"],
               [[s.get("size"), s.get("qty"), s.get("share")]
                for s in summary.get("sizeMix", [])])

    _add_sheet(wb, "Top add-ons", ["Add-on", "Servings"],
               [[a.get("name"), a.get("qty")] for a in summary.get("topAddons", [])])

    _add_sheet(wb, "Add-on affinity", ["Add-on A", "Add-on B", "Together"],
               [[a.get("a"), a.get("b"), a.get("count")] for a in summary.get("affinity", [])])

    by_hour = summary.get("ordersByHour") or {}
    _add_sheet(wb, "By hour", ["Hour", "Orders"],
               [[f"{h:02d}:00", by_hour.get(h, by_hour.get(str(h), 0))] for h in range(24)])
    return _save(wb)
